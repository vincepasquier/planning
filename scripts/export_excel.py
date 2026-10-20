#!/usr/bin/env python3
"""Exporte le planning en classeur Excel, et reimporte les durees corrigees.

Usage:
  python3 -I export_excel.py --debut 2026-10-26            # ecrit le classeur
  python3 -I export_excel.py --reimporter planning.xlsx    # relit l'onglet Durees

Le classeur comporte quatre onglets :
  Planning   le planning date, une ligne par lot et par local
  Durees     une ligne par tache, colonne Duree modifiable - c'est l'onglet
             a envoyer aux entreprises pour qu'elles corrigent les durees
  Questions  les points a trancher, par priorite
  Locaux     les locaux issus de la maquette, avec surfaces et hauteurs

La boucle de travail : on exporte, les entreprises corrigent la colonne Duree,
on reimporte, on recalcule.
"""
import argparse
import json
import sys
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from planifier import charger, planifier  # noqa: E402

ENTETE = PatternFill("solid", fgColor="1F3864")
ALERTE = PatternFill("solid", fgColor="FCE4E4")
JALON = PatternFill("solid", fgColor="FFF2CC")
BLANC = Font(color="FFFFFF", bold=True)
FIN = Side(style="thin", color="BFBFBF")
CADRE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
STATUTS_ALERTE = {"BLOQUE", "non_adjuge", "HORS_PERIMETRE_CONNU"}


def styler(ws, largeurs, nb_lignes):
    for i, l in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = l
    for c in ws[1]:
        c.fill, c.font = ENTETE, BLANC
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if nb_lignes:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(largeurs))}{nb_lignes + 1}"


def ecrire(args):
    taches, _ = charger(args.taches)
    a, m, j = map(int, args.debut.split("-"))
    dates = planifier(taches, date(a, m, j))
    wb = Workbook()

    # --- Planning
    ws = wb.active
    ws.title = "Planning"
    ws.append(["Lot", "Local", "Tâche", "Début", "Fin", "Sem.", "Phase", "Statut", "ID"])
    n = 0
    lignes = []
    for tid, (d0, f0) in dates.items():
        t = taches[tid]
        for local in (t["locaux"] or ["*"]):
            lignes.append((d0, f0, t, local, tid))
    for d0, f0, t, local, tid in sorted(lignes, key=lambda r: (r[0], r[2]["entreprise"])):
        ws.append([t["entreprise"], local if local != "*" else "(tous locaux)",
                   t["designation"], d0, f0, t.get("duree_semaines"),
                   t["phase"], t["statut"], tid])
        n += 1
        for c in ws[n + 1]:
            c.border = CADRE
            if t["statut"] in STATUTS_ALERTE:
                c.fill = ALERTE
            elif t.get("date_connue"):
                c.fill = JALON
        for col in ("D", "E"):
            ws[f"{col}{n + 1}"].number_format = "DD.MM.YYYY"
    styler(ws, [16, 18, 52, 12, 12, 7, 8, 14, 12], n)

    # --- Durees : l'onglet de travail
    ws = wb.create_sheet("Durees")
    ws.append(["ID", "Lot", "Tâche", "Locaux", "Durée (semaines)",
               "Durée corrigée", "Commentaire entreprise"])
    for i, (tid, t) in enumerate(sorted(taches.items(),
                                        key=lambda kv: (kv[1]["entreprise"], kv[0])), start=2):
        ws.append([tid, t["entreprise"], t["designation"], ", ".join(t["locaux"]),
                   t.get("duree_semaines"), None, None])
        for c in ws[i]:
            c.border = CADRE
        ws[f"F{i}"].fill = PatternFill("solid", fgColor="E2EFDA")
        ws[f"G{i}"].fill = PatternFill("solid", fgColor="E2EFDA")
    styler(ws, [12, 16, 52, 30, 16, 16, 44], len(taches))
    ws["F1"].comment = None

    # --- Questions
    ws = wb.create_sheet("Questions")
    ws.append(["ID", "Priorité", "Sujet", "Question", "Pourquoi", "Destinataire", "Impact"])
    q = json.load(open("data/questions-ouvertes.json", encoding="utf-8"))["questions"]
    rang = {"critique": 0, "haute": 1, "moyenne": 2, "basse": 3, "resolue": 4}
    for i, x in enumerate(sorted(q, key=lambda x: rang.get(x["priorite"], 9)), start=2):
        ws.append([x["id"], x["priorite"], x["sujet"], x["question"],
                   x.get("pourquoi") or x.get("reponse", ""),
                   x.get("destinataire", ""), x.get("impact", "")])
        for c in ws[i]:
            c.border = CADRE
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if x["priorite"] == "critique":
                c.fill = ALERTE
    styler(ws, [9, 11, 30, 56, 56, 26, 40], len(q))

    # --- Locaux
    ws = wb.create_sheet("Locaux")
    ws.append(["Niveau", "Numéro", "Désignation", "Surface nette m2", "HSP m", "Code"])
    loc = [x for x in json.load(open("data/extracted/locaux.json",
                                     encoding="utf-8"))["locaux"] if x["est_local"]]
    for i, x in enumerate(loc, start=2):
        ws.append([x["niveau"], x["numero"], x["designation"],
                   x.get("surface_nette_m2"), x.get("hauteur_sous_plafond_m"),
                   x.get("code_surface")])
        for c in ws[i]:
            c.border = CADRE
    styler(ws, [10, 16, 30, 17, 10, 8], len(loc))

    wb.save(args.sortie)
    print(f"{n} lignes de planning, {len(taches)} durées, {len(q)} questions, "
          f"{len(loc)} locaux -> {args.sortie}")


def reimporter(args):
    """Relit la colonne 'Duree corrigee' et met a jour taches.json."""
    wb = load_workbook(args.reimporter, data_only=True)
    if "Durees" not in wb.sheetnames:
        sys.exit("Onglet 'Durees' absent du classeur.")
    ws = wb["Durees"]
    corrections = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        tid, _, _, _, actuelle, corrigee = row[:6]
        if tid and corrigee is not None:
            try:
                v = int(round(float(corrigee)))
            except (TypeError, ValueError):
                continue
            if v > 0 and v != actuelle:
                corrections[tid] = v
    if not corrections:
        print("Aucune durée corrigée trouvée.")
        return
    data = json.load(open(args.taches, encoding="utf-8"))
    for t in data["taches"]:
        if t["id"] in corrections:
            print(f"  {t['id']:<12} {t.get('duree_semaines')} -> {corrections[t['id']]} semaines")
            t["duree_semaines"] = corrections[t["id"]]
    json.dump(data, open(args.taches, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"{len(corrections)} durée(s) mise(s) à jour dans {args.taches}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--taches", default="data/taches.json")
    ap.add_argument("--debut", default=date.today().isoformat())
    ap.add_argument("--sortie", default="data/extracted/planning.xlsx")
    ap.add_argument("--reimporter", metavar="FICHIER.xlsx")
    args = ap.parse_args()
    reimporter(args) if args.reimporter else ecrire(args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
